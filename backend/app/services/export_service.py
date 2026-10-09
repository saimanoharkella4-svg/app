import io
from datetime import datetime, timezone
from typing import List, Dict, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle


def generate_excel_report(daily_records: List[Dict[str, Any]], weekly_records: List[Dict[str, Any]]) -> io.BytesIO:
    """
    Generate a formatted multi-sheet Excel workbook with daily and weekly performance data.
    """
    wb = openpyxl.Workbook()
    
    # ------------------ SHEET 1: Daily Report ------------------
    ws_daily = wb.active
    ws_daily.title = "Daily Staff Report"
    
    # Styles
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10)
    alt_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    daily_headers = [
        "Employee Code", "Staff Name", "Date", "Login Time", "Logout Time", 
        "Working Hours", "Distance (KM)", "GPS Points", "Jobs Assigned", "Jobs Completed"
    ]
    
    # Write Title
    ws_daily.merge_cells("A1:J1")
    title_cell = ws_daily["A1"]
    title_cell.value = "CogniTrack Field Staff Tracker - Daily Summary Report"
    title_cell.font = Font(name="Segoe UI", size=14, bold=True, color="0F172A")
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    ws_daily.row_dimensions[1].height = 30

    # Write Headers
    ws_daily.append([]) # row 2 empty
    ws_daily.append(daily_headers) # row 3
    ws_daily.row_dimensions[3].height = 25

    for col_idx in range(1, len(daily_headers) + 1):
        cell = ws_daily.cell(row=3, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    # Write Data Rows
    row_idx = 4
    for item in daily_records:
        hours = item.get("working_minutes", 0) // 60
        mins = item.get("working_minutes", 0) % 60
        duration_str = f"{hours}h {mins:02d}m"
        
        row_values = [
            item.get("employee_code", ""),
            item.get("staff_name", ""),
            str(item.get("date", "")),
            item.get("login_time", "--"),
            item.get("logout_time", "--"),
            duration_str,
            round(item.get("distance_km", 0.0), 2),
            item.get("location_count", 0),
            item.get("jobs_assigned", 0),
            item.get("jobs_completed", 0)
        ]
        ws_daily.append(row_values)
        
        for col_idx in range(1, len(row_values) + 1):
            cell = ws_daily.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center" if col_idx in [1, 3, 4, 5, 6, 8, 9, 10] else "left")
            if row_idx % 2 == 1:
                cell.fill = alt_fill
        row_idx += 1

    # Adjust column widths
    for col in ws_daily.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws_daily.column_dimensions[col_letter].width = max(max_len + 3, 14)

    # ------------------ SHEET 2: Weekly Report ------------------
    ws_weekly = wb.create_sheet(title="Weekly Staff Report")
    
    # Title
    ws_weekly.merge_cells("A1:J1")
    title_weekly = ws_weekly["A1"]
    title_weekly.value = "CogniTrack Field Staff Tracker - Weekly Performance Report"
    title_weekly.font = Font(name="Segoe UI", size=14, bold=True, color="0F172A")
    title_weekly.alignment = Alignment(horizontal="left", vertical="center")
    ws_weekly.row_dimensions[1].height = 30

    weekly_headers = [
        "Employee Code", "Staff Name", "Week Start", "Week End", "Days Worked",
        "Total Hours", "Total Distance (KM)", "Avg Daily Hours", "Avg Daily Dist (KM)", "Jobs Completed"
    ]
    ws_weekly.append([])
    ws_weekly.append(weekly_headers)
    ws_weekly.row_dimensions[3].height = 25

    for col_idx in range(1, len(weekly_headers) + 1):
        cell = ws_weekly.cell(row=3, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    row_idx = 4
    for item in weekly_records:
        row_values = [
            item.get("employee_code", ""),
            item.get("staff_name", ""),
            str(item.get("week_start", "")),
            str(item.get("week_end", "")),
            item.get("days_worked", 0),
            round(item.get("total_hours", 0.0), 1),
            round(item.get("total_distance", 0.0), 2),
            round(item.get("average_daily_hours", 0.0), 1),
            round(item.get("average_daily_distance", 0.0), 2),
            item.get("jobs_completed", 0)
        ]
        ws_weekly.append(row_values)
        for col_idx in range(1, len(row_values) + 1):
            cell = ws_weekly.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center" if col_idx in [1, 3, 4, 5, 6, 8, 9, 10] else "left")
            if row_idx % 2 == 1:
                cell.fill = alt_fill
        row_idx += 1

    for col in ws_weekly.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws_weekly.column_dimensions[col_letter].width = max(max_len + 3, 14)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def generate_pdf_report(records: List[Dict[str, Any]], title: str = "CogniTrack FST Staff Report") -> io.BytesIO:
    """
    Generate an executive PDF report with headers, metrics, and formatted table.
    """
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(letter),
        rightMargin=30,
        leftMargin=30,
        topMargin=30,
        bottomMargin=30
    )

    styles = getSampleStyleSheet()
    
    # Custom styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=18,
        leading=22,
        textColor=colors.HexColor('#0F172A'),
        fontName='Helvetica-Bold'
    )
    subtitle_style = ParagraphStyle(
        'DocSub',
        parent=styles['Normal'],
        fontSize=10,
        textColor=colors.HexColor('#64748B'),
        fontName='Helvetica'
    )
    th_style = ParagraphStyle(
        'TableHeader',
        fontSize=9,
        leading=11,
        textColor=colors.white,
        fontName='Helvetica-Bold',
        alignment=1 # Center
    )
    td_style = ParagraphStyle(
        'TableCell',
        fontSize=8,
        leading=10,
        textColor=colors.HexColor('#1E293B'),
        fontName='Helvetica',
        alignment=1 # Center
    )
    td_left = ParagraphStyle(
        'TableCellLeft',
        fontSize=8,
        leading=10,
        textColor=colors.HexColor('#1E293B'),
        fontName='Helvetica',
        alignment=0 # Left
    )

    elements = []

    # Header
    elements.append(Paragraph(f"<b>CogniTrack</b> — Field Staff Tracker (FST)", title_style))
    generated_at = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    elements.append(Paragraph(f"{title} | Generated: {generated_at} | Internal Confidential", subtitle_style))
    elements.append(Spacer(1, 15))

    # Table Header
    headers = [
        "Code", "Name", "Date / Period", "Login", "Logout",
        "Hours", "Distance", "Points", "Jobs"
    ]
    table_data = [[Paragraph(h, th_style) for h in headers]]

    for r in records:
        hours_str = r.get("working_hours_formatted") or f"{r.get('total_hours', 0)}h"
        dist_str = f"{r.get('distance_km', r.get('total_distance', 0.0)):.1f} km"
        date_str = str(r.get("date") or f"{r.get('week_start')} ~ {r.get('week_end')}")
        
        row = [
            Paragraph(str(r.get("employee_code", "")), td_style),
            Paragraph(str(r.get("staff_name", "")), td_left),
            Paragraph(date_str, td_style),
            Paragraph(str(r.get("login_time", "--"))[:16], td_style),
            Paragraph(str(r.get("logout_time", "--"))[:16], td_style),
            Paragraph(hours_str, td_style),
            Paragraph(dist_str, td_style),
            Paragraph(str(r.get("location_count", "--")), td_style),
            Paragraph(f"{r.get('jobs_completed', 0)}/{r.get('jobs_assigned', 0)}", td_style),
        ]
        table_data.append(row)

    col_widths = [60, 110, 110, 85, 85, 70, 75, 55, 65]
    t = Table(table_data, colWidths=col_widths, repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E293B')),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
        ('TOPPADDING', (0, 0), (-1, 0), 6),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ('TOPPADDING', (0, 1), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 1), (-1, -1), 5),
    ]))
    elements.append(t)
    elements.append(Spacer(1, 15))
    elements.append(Paragraph("End of Report — CogniTrack Field Operations", subtitle_style))

    doc.build(elements)
    buffer.seek(0)
    return buffer


def generate_individual_weekly_pdf_report(
    staff_name: str,
    employee_code: str,
    department: str,
    week_start: str,
    week_end: str,
    weekly_totals: Dict[str, Any],
    daily_breakdown: List[Dict[str, Any]]
) -> io.BytesIO:
    """
    Generate a styled PDF report specifically for an individual monitored staff member for a 7-day week.
    """
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(letter),
        rightMargin=30,
        leftMargin=30,
        topMargin=30,
        bottomMargin=30
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('DocTitle', parent=styles['Heading1'], fontSize=18, leading=22, textColor=colors.HexColor('#0F172A'), fontName='Helvetica-Bold')
    subtitle_style = ParagraphStyle('DocSub', parent=styles['Normal'], fontSize=10, textColor=colors.HexColor('#64748B'), fontName='Helvetica')
    label_style = ParagraphStyle('LabelStyle', fontSize=10, textColor=colors.HexColor('#475569'), fontName='Helvetica-Bold')
    val_style = ParagraphStyle('ValStyle', fontSize=10, textColor=colors.HexColor('#0F172A'), fontName='Helvetica')
    th_style = ParagraphStyle('TableHeader', fontSize=9, leading=11, textColor=colors.white, fontName='Helvetica-Bold', alignment=1)
    td_style = ParagraphStyle('TableCell', fontSize=8, leading=10, textColor=colors.HexColor('#1E293B'), fontName='Helvetica', alignment=1)

    elements = []
    elements.append(Paragraph("<b>CogniTrack</b> — Field Operations Telemetry Console", title_style))
    generated_at = datetime.now(timezone.utc).strftime("%B %d, %Y %H:%M UTC")
    elements.append(Paragraph(f"Executive Weekly Operations Report | Generated: {generated_at}", subtitle_style))
    elements.append(Spacer(1, 14))

    # Executive & Weekly Totals Grid
    meta_data = [
        [Paragraph("Executive Name:", label_style), Paragraph(staff_name, val_style), Paragraph("Week Window:", label_style), Paragraph(f"{week_start} to {week_end}", val_style)],
        [Paragraph("Employee Code:", label_style), Paragraph(employee_code, val_style), Paragraph("Days Worked:", label_style), Paragraph(f"{weekly_totals.get('days_worked', 0)} Days", val_style)],
        [Paragraph("Department:", label_style), Paragraph(department or "Field Operations", val_style), Paragraph("Total Hours:", label_style), Paragraph(str(weekly_totals.get('total_hours_formatted', '0h 0m')), val_style)],
        [Paragraph("Total Distance:", label_style), Paragraph(f"{weekly_totals.get('total_distance', 0):.1f} KM", val_style), Paragraph("Tasks Completed:", label_style), Paragraph(f"{weekly_totals.get('jobs_completed', 0)} Tasks", val_style)],
        [Paragraph("Avg Daily Dist:", label_style), Paragraph(f"{weekly_totals.get('average_daily_distance', 0):.1f} KM/day", val_style), Paragraph("Avg Daily Hours:", label_style), Paragraph(f"{weekly_totals.get('average_daily_hours', 0):.1f} hrs/day", val_style)]
    ]

    t_meta = Table(meta_data, colWidths=[110, 240, 110, 240])
    t_meta.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F8FAFC')),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ('PADDING', (0, 0), (-1, -1), 6),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    elements.append(t_meta)
    elements.append(Spacer(1, 16))

    elements.append(Paragraph("<b>7-Day Daily Operations Breakdown</b>", title_style))
    elements.append(Spacer(1, 8))

    headers = ["Day", "Date", "Working Hours", "Distance (KM)", "Tasks Completed"]
    table_data = [[Paragraph(h, th_style) for h in headers]]

    for d in daily_breakdown:
        table_data.append([
            Paragraph(str(d.get("day_name", "")), td_style),
            Paragraph(str(d.get("date", "")), td_style),
            Paragraph(str(d.get("working_hours_formatted", "0h 0m")), td_style),
            Paragraph(f"{d.get('distance_km', 0.0):.1f} KM", td_style),
            Paragraph(f"{d.get('jobs_completed', 0)}", td_style)
        ])

    col_widths = [140, 140, 140, 140, 140]
    t_daily = Table(table_data, colWidths=col_widths, repeatRows=1)
    t_daily.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4F46E5')),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ('PADDING', (0, 0), (-1, -1), 6),
    ]))
    elements.append(t_daily)
    elements.append(Spacer(1, 18))
    elements.append(Paragraph("End of Individual Executive Report — CogniTrack Fleet Telemetry Platform", subtitle_style))

    doc.build(elements)
    buffer.seek(0)
    return buffer
